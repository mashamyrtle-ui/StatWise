"""
Document processing service.
Extracts text from PDF, DOCX and PPTX files and chunks it for AI processing.
"""
import os
import re
import logging
from typing import List, Tuple, Optional

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> Tuple[str, int]:
    """Extract text from a PDF file using pypdf. Returns (text, page_count)."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        pages = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[Page {i}]\n{text}")
        return "\n\n".join(pages), len(reader.pages)
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        return "", 0


def extract_text_from_docx(file_path: str) -> Tuple[str, int]:
    """Extract text from a DOCX file. Returns (text, paragraph_count)."""
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs), len(paragraphs)
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        return "", 0


def extract_text_from_pptx(file_path: str) -> Tuple[str, int]:
    """Extract text from a PPTX file. Returns (text, slide_count)."""
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        slides = []
        for slide_num, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides.append(f"[Slide {slide_num}]\n" + "\n".join(texts))
        return "\n\n".join(slides), len(prs.slides)
    except Exception as e:
        logger.error(f"PPTX extraction error: {e}")
        return "", 0


def extract_text(file_path: str, file_type: str) -> Tuple[str, int]:
    """Route to the correct extractor based on file type."""
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type == "docx":
        return extract_text_from_docx(file_path)
    elif file_type == "pptx":
        return extract_text_from_pptx(file_path)
    else:
        return "", 0


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 100) -> List[str]:
    """Split text into overlapping chunks for AI processing."""
    if not text:
        return []
    words = text.split()
    chunks = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunk = " ".join(words[start:end])
        if chunk.strip():
            chunks.append(chunk)
        start += chunk_size - overlap
    return chunks


def detect_concepts(text: str) -> List[str]:
    """Detect statistical concepts mentioned in the text."""
    concept_keywords = {
        "sampling": ["sampling", "sample", "population", "stratum", "cluster", "pps"],
        "statistics": ["mean", "median", "variance", "standard deviation", "hypothesis", "regression"],
        "survey design": ["questionnaire", "survey", "fieldwork", "enumeration", "response rate"],
        "python": ["python", "pandas", "numpy", "matplotlib", "scikit"],
        "sql": ["sql", "database", "query", "join", "select"],
        "gis": ["gis", "spatial", "geographic", "coordinate", "shapefile", "raster"],
        "machine learning": ["machine learning", "classification", "clustering", "neural", "model"],
        "data visualization": ["visualization", "chart", "graph", "plot", "dashboard"],
        "national accounts": ["gdp", "gnp", "national accounts", "input-output"],
        "price statistics": ["cpi", "wpi", "inflation", "price index"],
    }

    text_lower = text.lower()
    detected = []
    for concept, keywords in concept_keywords.items():
        if any(kw in text_lower for kw in keywords):
            detected.append(concept)
    return detected


def detect_competencies(text: str) -> List[str]:
    """Map detected concepts to formal competency names."""
    concept_to_competency = {
        "sampling": "Sampling",
        "statistics": "Statistical Analysis",
        "survey design": "Survey Design",
        "python": "Python",
        "sql": "SQL",
        "gis": "GIS",
        "machine learning": "AI/ML",
        "data visualization": "Data Visualization",
        "national accounts": "National Accounts",
        "price statistics": "Price Statistics",
    }
    concepts = detect_concepts(text)
    return [concept_to_competency[c] for c in concepts if c in concept_to_competency]


def get_allowed_extensions() -> List[str]:
    return ["pdf", "docx", "pptx"]


def validate_file(filename: str, file_size: int, max_size_mb: int) -> Optional[str]:
    """Return error message if file is invalid, else None."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in get_allowed_extensions():
        return f"File type .{ext} is not supported. Please upload PDF, DOCX, or PPTX files."
    if file_size > max_size_mb * 1024 * 1024:
        return f"File size exceeds the {max_size_mb}MB limit."
    return None
